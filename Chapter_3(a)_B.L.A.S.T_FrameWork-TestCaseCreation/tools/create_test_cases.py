"""
Layer 3 Tool: Create test cases in Excel + XML format from image data.
"""
import os
import xml.etree.ElementTree as ET
from xml.dom import minidom
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)

_ROOT   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = _ROOT

# ── Test Case Data ──────────────────────────────────────────────
TEST_CASES = [
    {
        "test_case":          "Test_01",
        "tax_plan_code":      "1247Q",
        "plan_booster_entity":"NOT PRVCX and\nNOT PRVQX and\nNOT PRVZX and\nNOT PRVTX",
        "employer_plan_key":  "NOT 721450-X2\nNOT 721451-X2\nNOT 884203-X2\nNOT 664912-X2",
        "expected_result":    "New C12.0 - Green text and Light Black text should print - IN",
        "outcome":            "IN",
    },
    {
        "test_case":          "Test_02",
        "tax_plan_code":      "1247Q",
        "plan_booster_entity":"PRVCX",
        "employer_plan_key":  "NOT 721450-X2",
        "expected_result":    "New C12.0 - Green text and Light Black text should NOT print - OUT",
        "outcome":            "OUT",
    },
    {
        "test_case":          "Test_03",
        "tax_plan_code":      "1247Q",
        "plan_booster_entity":"NOT PRVCX",
        "employer_plan_key":  "721450-X2",
        "expected_result":    "New C12.0 - Green text and Light Black text should NOT print - OUT",
        "outcome":            "OUT",
    },
]

HEADERS = ["Test_Case", "TaxPlanCode", "Plan Booster Entity",
           "Employer Plan Key", "Expected Result"]

# ── Styles ──────────────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="1F3864")   # dark navy
IN_FILL      = PatternFill("solid", fgColor="E2EFDA")   # light green
OUT_FILL     = PatternFill("solid", fgColor="FCE4D6")   # light red/salmon
ALT_FILL     = PatternFill("solid", fgColor="F2F2F2")   # light grey
WHITE_FILL   = PatternFill("solid", fgColor="FFFFFF")

HEADER_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT    = Font(name="Calibri", size=10)
BOLD_FONT    = Font(name="Calibri", bold=True, size=10)

THIN  = Side(style="thin",   color="BFBFBF")
THICK = Side(style="medium", color="1F3864")
BORDER_THIN  = Border(left=THIN,  right=THIN,  top=THIN,  bottom=THIN)
BORDER_THICK = Border(left=THICK, right=THICK, top=THICK, bottom=THICK)

CENTER  = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT    = Alignment(horizontal="left",   vertical="center", wrap_text=True)

COL_WIDTHS = [12, 14, 30, 30, 50]


def build_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "KAN-5 Test Cases"

    # ── Title row ──
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value     = "KAN-5 — Muco TaxPlanCode 1247Q Enhancement · Test Cases"
    title_cell.font      = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
    title_cell.fill      = PatternFill("solid", fgColor="0052CC")
    title_cell.alignment = CENTER
    ws.row_dimensions[1].height = 28

    # ── Header row ──
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = CENTER
        cell.border    = BORDER_THIN
    ws.row_dimensions[2].height = 22

    # ── Data rows ──
    for row_idx, tc in enumerate(TEST_CASES, start=3):
        fill = IN_FILL if tc["outcome"] == "IN" else OUT_FILL

        values = [
            tc["test_case"],
            tc["tax_plan_code"],
            tc["plan_booster_entity"],
            tc["employer_plan_key"],
            tc["expected_result"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font      = BOLD_FONT if col_idx == 1 else BODY_FONT
            cell.fill      = fill
            cell.alignment = CENTER if col_idx <= 2 else LEFT
            cell.border    = BORDER_THIN
        ws.row_dimensions[row_idx].height = 60

    # ── Column widths ──
    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    # ── Freeze header rows ──
    ws.freeze_panes = "A3"

    out_path = os.path.join(OUT_DIR, "KAN5_TestCases.xlsx")
    wb.save(out_path)
    print(f"✅ Excel created: {out_path}")
    return out_path


def build_xml():
    root = ET.Element("TestCases")
    root.set("jira_issue", "KAN-5")
    root.set("summary",    "Muco TaxPlanCode 1247Q Enhancement")

    for tc in TEST_CASES:
        tc_el = ET.SubElement(root, "TestCase")
        tc_el.set("id", tc["test_case"])
        tc_el.set("outcome", tc["outcome"])

        ET.SubElement(tc_el, "TaxPlanCode").text      = tc["tax_plan_code"]
        ET.SubElement(tc_el, "PlanBoosterEntity").text = tc["plan_booster_entity"]
        ET.SubElement(tc_el, "EmployerPlanKey").text   = tc["employer_plan_key"]
        ET.SubElement(tc_el, "ExpectedResult").text    = tc["expected_result"]

    # Pretty-print
    raw    = ET.tostring(root, encoding="unicode")
    pretty = minidom.parseString(raw).toprettyxml(indent="  ")
    # Remove the extra xml declaration added by toprettyxml
    lines  = pretty.split("\n")[1:]
    pretty = '<?xml version="1.0" encoding="UTF-8"?>\n' + "\n".join(lines)

    out_path = os.path.join(OUT_DIR, "KAN5_TestCases.xml")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(pretty)
    print(f"✅ XML   created: {out_path}")
    return out_path


if __name__ == "__main__":
    build_excel()
    build_xml()
